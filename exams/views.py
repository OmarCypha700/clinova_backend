import csv

from django.db import transaction
from django.db.models import Count, OuterRef, Prefetch, Q, Subquery, Sum, Value
from django.db.models.functions import Coalesce
from django.http import HttpResponse
from django.utils import timezone
from django_filters.rest_framework import DjangoFilterBackend
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from rest_framework import filters, status, viewsets
from rest_framework.decorators import action
from rest_framework.generics import (
    ListAPIView,
    ListCreateAPIView,
    RetrieveAPIView,
    RetrieveUpdateDestroyAPIView,
)
from rest_framework.pagination import PageNumberPagination
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from accounts.models import User

from .filters import StudentFilter
from .models import (
    CarePlan,
    Level,
    Procedure,
    ProcedureStep,
    ProcedureStepScore,
    Program,
    ReconciledScore,
    SiteSettings,
    Student,
    StudentProcedure,
)
from .permissions import IsAdmin, IsExaminer
from .serializers import (
    CarePlanCreateSerializer,
    CarePlanSerializer,
    DashboardStatsSerializer,
    LevelSerializer,
    ProcedureAdminListSerializer,
    ProcedureCreateUpdateSerializer,
    ProcedureDetailSerializer,
    ProcedureListSerializer,
    ProcedureStepCreateUpdateSerializer,
    ProgramSerializer,
    ReconciliationSerializer,
    SiteSettingsSerializer,
    StudentCreateUpdateSerializer,
    StudentSerializer,
    UserCreateSerializer,
    UserSerializer,
)


# ─────────────────────────────────────────────
# PAGINATION CLASSES
# ─────────────────────────────────────────────


class StudentPagination(PageNumberPagination):
    page_size = 100
    page_size_query_param = "page_size"
    max_page_size = 5000


class GradesPagination(PageNumberPagination):
    page_size = 50
    page_size_query_param = "page_size"
    max_page_size = 500


class ProcedurePagination(PageNumberPagination):
    page_size = 100
    page_size_query_param = "page_size"
    max_page_size = 5000


# ─────────────────────────────────────────────
# SITESETTINGS VIEW
# ─────────────────────────────────────────────


class SiteSettingsView(APIView):
    """
    GET  /exams/settings/   – retrieve current settings
    PATCH /exams/settings/  – update one or more flags (admin only)
    """

    def get_permissions(self):
        if self.request.method == "GET":
            return [IsAuthenticated()]
        return [IsAuthenticated(), IsAdmin()]

    def get(self, request):
        settings = SiteSettings.get()
        return Response(SiteSettingsSerializer(settings).data)

    def patch(self, request):
        settings = SiteSettings.get()
        serializer = SiteSettingsSerializer(settings, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save(updated_by=request.user)
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


# ===============================================
# GENERAL VIEWS
# ===============================================

# ─────────────────────────────────────────────
# PROGRAM VIEWS
# ─────────────────────────────────────────────


class ProgramListView(ListAPIView):
    permission_classes = [IsAuthenticated, IsExaminer | IsAdmin]
    queryset = Program.objects.all()
    serializer_class = ProgramSerializer


class ProgramViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated, IsAdmin]
    serializer_class = ProgramSerializer

    def get_queryset(self):
        return Program.objects.annotate(
            student_count=Count("students", distinct=True),
            procedure_count=Count("procedures", distinct=True),
        ).order_by("id")


# ─────────────────────────────────────────────
# LEVEL VIEWS
# ─────────────────────────────────────────────


class LevelListCreateView(ListCreateAPIView):
    queryset = Level.objects.all()
    serializer_class = LevelSerializer
    permission_classes = [IsAuthenticated]


class LevelDetailView(RetrieveUpdateDestroyAPIView):
    queryset = Level.objects.all()
    serializer_class = LevelSerializer
    permission_classes = [IsAuthenticated]


# ===============================================
# EAMINER FACING VIEWS
# ===============================================

# ─────────────────────────────────────────────
# STUDENT VIEWS (EF)
# ─────────────────────────────────────────────


class StudentByProgramView(ListAPIView):
    permission_classes = [IsAuthenticated, IsExaminer | IsAdmin]
    serializer_class = StudentSerializer
    pagination_class = StudentPagination
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ["full_name", "index_number"]
    ordering_fields = ["index_number", "full_name"]
    ordering = ["index_number"]

    def get_queryset(self):
        program_id = self.kwargs["program_id"]
        queryset = Student.objects.select_related("program", "level").filter(
            program_id=program_id, is_active=True
        )
        level_id = self.request.query_params.get("level")
        if level_id and level_id != "all":
            queryset = queryset.filter(level_id=level_id)
        return queryset


class StudentDetailView(RetrieveAPIView):
    permission_classes = [IsAuthenticated, IsExaminer]
    queryset = Student.objects.select_related("program", "level")
    serializer_class = StudentSerializer


# ─────────────────────────────────────────────
# CARE PLAN (EF)
# ─────────────────────────────────────────────


class CarePlanView(APIView):
    """
    GET  – return existing care plan or {exists: False}
    POST – submit (or rescore, depending on the care_plan_lock_on_submit flag)
    """

    permission_classes = [IsAuthenticated]

    def get(self, request, student_id, program_id):
        try:
            care_plan = CarePlan.objects.select_related("examiner").get(
                student_id=student_id,
                program_id=program_id,
            )
            return Response(CarePlanSerializer(care_plan).data)
        except CarePlan.DoesNotExist:
            return Response(
                {"exists": False, "message": "No care plan found for this student"},
                status=status.HTTP_200_OK,
            )

    @transaction.atomic
    def post(self, request, student_id, program_id):
        site_settings = SiteSettings.get()
        lock_on_submit = site_settings.care_plan_lock_on_submit

        existing = CarePlan.objects.filter(
            student_id=student_id,
            program_id=program_id,
        ).first()

        if existing:
            if lock_on_submit:
                # Flag is ON (default) – prevent rescoring
                return Response(
                    {"error": "Care plan already submitted for this student"},
                    status=status.HTTP_400_BAD_REQUEST,
                )
            else:
                # Flag is OFF – allow rescoring: update in-place
                score = request.data.get("score")
                comments = request.data.get("comments", existing.comments)

                if score is None:
                    return Response(
                        {"error": "score is required"},
                        status=status.HTTP_400_BAD_REQUEST,
                    )

                score_int = int(score)
                if not (0 <= score_int <= 20):
                    return Response(
                        {"error": "Score must be between 0 and 20"},
                        status=status.HTTP_400_BAD_REQUEST,
                    )

                existing.score = score_int
                existing.comments = comments
                existing.examiner = request.user  # track who rescored
                existing.is_locked = False  # stays unlocked
                existing.save(
                    update_fields=["score", "comments", "examiner", "is_locked"]
                )
                return Response(
                    CarePlanSerializer(existing).data, status=status.HTTP_200_OK
                )

        # No existing care plan — create it
        data = request.data.copy()
        data["student"] = student_id
        data["program"] = program_id

        serializer = CarePlanCreateSerializer(data=data)
        if serializer.is_valid():
            care_plan = serializer.save(
                examiner=request.user,
                is_locked=lock_on_submit,  # lock only when flag is ON
            )
            return Response(
                CarePlanSerializer(care_plan).data,
                status=status.HTTP_201_CREATED,
            )
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


# Note: The original version of CarePlanView (below) was a simple create-once endpoint that locked the care plan after submission.
# The new version (above) adds support for optional rescoring based on a site setting, allowing updates to the care plan score and comments if the lock_on_submit flag is turned off.
# The original version is kept here for reference and potential rollback if needed.

# class CarePlanView(APIView):
#     permission_classes = [IsAuthenticated]

#     def get(self, request, student_id, program_id):
#         try:
#             care_plan = CarePlan.objects.select_related("student", "examiner").get(
#                 student_id=student_id, program_id=program_id
#             )
#             return Response(CarePlanSerializer(care_plan).data)
#         except CarePlan.DoesNotExist:
#             return Response({"exists": False, "message": "No care plan found"}, status=200)

#     @transaction.atomic
#     def post(self, request, student_id, program_id):
#         if CarePlan.objects.filter(student_id=student_id, program_id=program_id).exists():
#             return Response(
#                 {"error": "Care plan already submitted for this student"},
#                 status=status.HTTP_400_BAD_REQUEST,
#             )
#         data = {**request.data, "student": student_id, "program": program_id}
#         serializer = CarePlanCreateSerializer(data=data)
#         if serializer.is_valid():
#             care_plan = serializer.save(examiner=request.user, is_locked=True)
#             return Response(CarePlanSerializer(care_plan).data, status=201)
#         return Response(serializer.errors, status=400)


# ─────────────────────────────────────────────
# ASSIGN EXAMINERS (EF)
# ─────────────────────────────────────────────


class AssignExaminersView(APIView):
    permission_classes = [IsAuthenticated, IsExaminer]

    def post(self, request, *args, **kwargs):
        data = request.data
        required = ["student_id", "procedure_id", "examiner_a_id", "examiner_b_id"]
        if not all(data.get(f) for f in required):
            return Response({"detail": "All fields are required."}, status=400)
        try:
            student = Student.objects.get(id=data["student_id"])
            procedure = Procedure.objects.get(id=data["procedure_id"])
            examiner_a = User.objects.get(id=data["examiner_a_id"], role="examiner")
            examiner_b = User.objects.get(id=data["examiner_b_id"], role="examiner")
        except (Student.DoesNotExist, Procedure.DoesNotExist, User.DoesNotExist) as e:
            return Response({"detail": f"Invalid reference: {e}"}, status=400)
        sp, created = StudentProcedure.objects.update_or_create(
            student=student,
            procedure=procedure,
            defaults={"examiner_a": examiner_a, "examiner_b": examiner_b},
        )
        return Response(
            {
                "id": sp.id,
                "created": created,
                "examiner_a": examiner_a.get_full_name(),
                "examiner_b": examiner_b.get_full_name(),
            }
        )


# ─────────────────────────────────────────────
# PROCEDURE VIEWS (EF)
# ─────────────────────────────────────────────


# class ProcedureByProgramView(ListAPIView):
#     permission_classes = [IsAuthenticated, IsExaminer | IsAdmin]
#     serializer_class = ProcedureListSerializer
#     def get_queryset(self):
#         return Procedure.objects.filter(program_id=self.kwargs["program_id"]).annotate(
#             step_count=Count("steps")
#         )

#     def get_serializer_context(self):
#         context = super().get_serializer_context()
#         student_id = self.request.query_params.get("student_id")
#         context["student_id"] = student_id

#         # Pre-fetch all StudentProcedures for this student/program in ONE query
#         # and pass a lookup map to the serializer – eliminates N+1
#         if student_id:
#             sps = StudentProcedure.objects.filter(
#                 student_id=student_id,
#                 procedure__program_id=self.kwargs["program_id"],
#             ).select_related("examiner_a", "examiner_b", "assigned_reconciler")
#             context["student_procedures_map"] = {sp.procedure_id: sp for sp in sps}
#         return context


class ProcedureByProgramView(ListAPIView):
    permission_classes = [IsAuthenticated, IsExaminer | IsAdmin]
    pagination_class = ProcedurePagination
    filter_backends = [DjangoFilterBackend, filters.SearchFilter]
    search_fields = ["name"]
    serializer_class = ProcedureListSerializer

    def get_queryset(self):
        return (
            Procedure.objects.filter(program_id=self.kwargs["program_id"])
            .annotate(step_count=Count("steps"))
            .order_by("pk")
        )

    def get_serializer_context(self):
        context = super().get_serializer_context()

        student_id = self.request.query_params.get("student_id")

        try:
            student_id = int(student_id) if student_id else None
        except (TypeError, ValueError):
            student_id = None

        context["student_id"] = student_id

        if student_id:
            sps = StudentProcedure.objects.filter(
                student_id=student_id,
                procedure__program_id=self.kwargs["program_id"],
            ).select_related("examiner_a", "examiner_b", "assigned_reconciler")

            context["student_procedures_map"] = {sp.procedure_id: sp for sp in sps}

        return context


class ProcedureDetailView(RetrieveAPIView):
    permission_classes = [IsAuthenticated, IsExaminer]
    queryset = Procedure.objects.prefetch_related("steps")
    serializer_class = ProcedureDetailSerializer

    def retrieve(self, request, *args, **kwargs):
        student_id = self.kwargs.get("student_id")
        procedure = self.get_object()

        sp, _ = StudentProcedure.objects.select_related(
            "examiner_a", "examiner_b", "assigned_reconciler"
        ).get_or_create(
            student_id=student_id,
            procedure=procedure,
            defaults={
                "examiner_a": None,  # No assignment on open — deferred to first score save
                "examiner_b": None,
            },
        )

        both_slots_filled = (
            sp.examiner_a is not None
            and sp.examiner_b is not None
            and sp.examiner_a != sp.examiner_b
        )
        is_assigned = request.user in (sp.examiner_a, sp.examiner_b)

        # Block non-assigned examiners from viewing scored or reconciled procedures
        if sp.status in ("scored", "reconciled") and not is_assigned:
            return Response(
                {
                    "detail": "This procedure has been scored and is no longer accessible.",
                    "is_locked": True,
                },
                status=status.HTTP_403_FORBIDDEN,
            )

        # Block non-assigned examiners once both slots are claimed
        if both_slots_filled and not is_assigned:
            return Response(
                {
                    "detail": "You are not assigned as an examiner for this procedure.",
                    "examiner_a": sp.examiner_a.get_full_name(),
                    "examiner_b": sp.examiner_b.get_full_name(),
                    "is_locked": sp.assigned_reconciler is not None,
                },
                status=status.HTTP_403_FORBIDDEN,
            )

        # Block further scoring once a reconciler is assigned
        if sp.assigned_reconciler and sp.status != "reconciled":
            return Response(
                {
                    "detail": "This procedure is locked. A reconciler has been assigned.",
                    "assigned_reconciler": sp.assigned_reconciler.get_full_name(),
                    "is_locked": True,
                },
                status=status.HTTP_403_FORBIDDEN,
            )

        return super().retrieve(request, *args, **kwargs)

    def get_serializer_context(self):
        context = super().get_serializer_context()
        context["student_id"] = self.kwargs.get("student_id")
        return context


# ─────────────────────────────────────────────
# AUTOSAVE SCORE (EF)
# ─────────────────────────────────────────────


class AutosaveStepScoreView(APIView):
    permission_classes = [IsAuthenticated, IsExaminer]

    @transaction.atomic
    def post(self, request, *args, **kwargs):
        data = request.data
        student_procedure_id = data.get("student_procedure")
        step_id = data.get("step")
        score = data.get("score")

        if not all([student_procedure_id, step_id, score is not None]):
            return Response(
                {"detail": "student_procedure, step, and score are required."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            sp = (
                StudentProcedure.objects.select_related(
                    "examiner_a", "examiner_b", "assigned_reconciler", "procedure"
                )
                .select_for_update()
                .get(id=student_procedure_id)
            )
            step = ProcedureStep.objects.get(id=step_id, procedure=sp.procedure)
        except StudentProcedure.DoesNotExist:
            return Response({"detail": "StudentProcedure not found."}, status=404)
        except ProcedureStep.DoesNotExist:
            return Response({"detail": "ProcedureStep not found."}, status=404)

        # if request.user not in (sp.examiner_a, sp.examiner_b):
        #     return Response(
        #         {"detail": "You are not authorized to score this procedure."},
        #         status=status.HTTP_403_FORBIDDEN,
        #     )

        if request.user not in (sp.examiner_a, sp.examiner_b):
            # Slot A is still open — claim it now on first score save
            if sp.examiner_a is None:
                sp.examiner_a = request.user
                sp.save(update_fields=["examiner_a"])
            # Slot B is still open — claim it now on first score save
            elif sp.examiner_b is None or sp.examiner_a == sp.examiner_b:
                sp.examiner_b = request.user
                sp.save(update_fields=["examiner_b"])
            # Both slots genuinely taken — deny
            else:
                return Response(
                    {"detail": "You are not authorized to score this procedure."},
                    status=status.HTTP_403_FORBIDDEN,
                )

        if sp.assigned_reconciler:
            return Response(
                {"detail": "Cannot modify scores. Reconciler has been assigned."},
                status=status.HTTP_403_FORBIDDEN,
            )
        if sp.status == "reconciled":
            return Response(
                {"detail": "Cannot modify scores. Procedure has been reconciled."},
                status=status.HTTP_403_FORBIDDEN,
            )

        step_score, created = ProcedureStepScore.objects.update_or_create(
            student_procedure=sp,
            step=step,
            examiner=request.user,
            defaults={"score": score},
        )

        examiner_a_complete = False
        examiner_b_complete = False

        # if sp.examiner_a != sp.examiner_b:
        #     total_steps = sp.procedure.steps.count()
        #     score_map = {
        #         s["examiner"]: s["c"]
        #         for s in sp.step_scores.values("examiner").annotate(c=Count("id"))
        #     }
        #     examiner_a_complete = score_map.get(sp.examiner_a_id, 0) == total_steps
        #     examiner_b_complete = score_map.get(sp.examiner_b_id, 0) == total_steps

        #     if examiner_a_complete and examiner_b_complete and sp.status == "pending":
        #         sp.status = "scored"
        #         sp.save(update_fields=["status"])

        both_assigned = (
            sp.examiner_a is not None
            and sp.examiner_b is not None
            and sp.examiner_a != sp.examiner_b
        )
        if both_assigned:
            total_steps = sp.procedure.steps.count()
            score_map = {
                s["examiner"]: s["c"]
                for s in sp.step_scores.values("examiner").annotate(c=Count("id"))
            }
            examiner_a_complete = score_map.get(sp.examiner_a_id, 0) == total_steps
            examiner_b_complete = score_map.get(sp.examiner_b_id, 0) == total_steps

            if examiner_a_complete and examiner_b_complete and sp.status == "pending":
                sp.status = "scored"
                sp.save(update_fields=["status"])

        return Response(
            {
                "step": step.id,
                "score": step_score.score,
                "created": created,
                "status": sp.status,
                "examiner_a_complete": examiner_a_complete,
                "examiner_b_complete": examiner_b_complete,
                "both_examiners_assigned": both_assigned,
                "is_locked": sp.assigned_reconciler is not None,
            },
            status=status.HTTP_200_OK,
        )


# ─────────────────────────────────────────────
# RECONCILIATION (EF)
# ─────────────────────────────────────────────


class ReconciliationView(RetrieveAPIView):
    permission_classes = [IsAuthenticated, IsExaminer]
    serializer_class = ReconciliationSerializer

    def get_queryset(self):
        return StudentProcedure.objects.filter(
            student_id=self.kwargs["student_id"],
            procedure_id=self.kwargs["procedure_id"],
        )

    def get_object(self):
        obj = (
            self.get_queryset()
            .select_related(
                "examiner_a",
                "examiner_b",
                "reconciled_by",
                "assigned_reconciler",
                "student",
                "procedure",
            )
            .prefetch_related(
                Prefetch(
                    "step_scores",
                    queryset=ProcedureStepScore.objects.select_related("examiner"),
                ),
                Prefetch(
                    "reconciled_scores",
                    queryset=ReconciledScore.objects.select_related("step"),
                ),
                "procedure__steps",
            )
            .first()
        )

        if not obj:
            obj = StudentProcedure.objects.create(
                student_id=self.kwargs["student_id"],
                procedure_id=self.kwargs["procedure_id"],
                examiner_a=self.request.user,
                examiner_b=self.request.user,
            )

        if obj.status == "scored" and not obj.assigned_reconciler:
            if obj.can_user_reconcile(self.request.user):
                obj.assigned_reconciler = self.request.user
                obj.save(update_fields=["assigned_reconciler"])

        return obj


class SaveReconciliationView(APIView):
    permission_classes = [IsAuthenticated, IsExaminer]

    @transaction.atomic
    def post(self, request, *args, **kwargs):
        student_procedure_id = request.data.get("student_procedure_id")
        reconciled_scores = request.data.get("reconciled_scores", [])

        if not student_procedure_id or not reconciled_scores:
            return Response(
                {"detail": "student_procedure_id and reconciled_scores are required."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            sp = (
                StudentProcedure.objects.select_related("procedure")
                .select_for_update()
                .get(id=student_procedure_id)
            )
        except StudentProcedure.DoesNotExist:
            return Response({"detail": "StudentProcedure not found."}, status=404)

        if not sp.can_user_reconcile(request.user):
            return Response(
                {"detail": "Not authorized to reconcile this procedure."},
                status=status.HTTP_403_FORBIDDEN,
            )

        total_steps = sp.procedure.steps.count()
        if len(reconciled_scores) != total_steps:
            return Response(
                {
                    "detail": f"Expected {total_steps} scores, got {len(reconciled_scores)}."
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Validate all step IDs up-front before any writes
        step_ids = [s.get("step_id") for s in reconciled_scores]
        if None in step_ids:
            return Response(
                {"detail": "Each score must have step_id and score."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        valid_steps = {
            s.id: s
            for s in ProcedureStep.objects.filter(
                id__in=step_ids, procedure=sp.procedure
            )
        }
        if len(valid_steps) != total_steps:
            return Response(
                {"detail": "One or more step IDs are invalid for this procedure."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        # Replace reconciled scores atomically
        sp.reconciled_scores.all().delete()
        ReconciledScore.objects.bulk_create(
            [
                ReconciledScore(
                    student_procedure=sp,
                    step=valid_steps[s["step_id"]],
                    score=s["score"],
                    reconciled_by=request.user,
                )
                for s in reconciled_scores
            ]
        )

        sp.status = "reconciled"
        sp.reconciled_by = request.user
        sp.reconciled_at = timezone.now()
        sp.save(update_fields=["status", "reconciled_by", "reconciled_at"])

        return Response(
            {
                "detail": "Reconciliation saved successfully.",
                "status": sp.status,
                "reconciled_by": request.user.get_full_name(),
                "reconciled_at": sp.reconciled_at,
            },
            status=status.HTTP_200_OK,
        )


# ===============================================
# ADMIN FACING VIEWS
# ===============================================

# ─────────────────────────────────────────────
# ADMIN: DASHBOARD
# ─────────────────────────────────────────────


class DashboardStatsView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        from django.db.models import Count as _Count

        student_agg = Student.objects.aggregate(
            total=_Count("id"),
            active=_Count("id", filter=Q(is_active=True)),
        )
        sp_agg = StudentProcedure.objects.aggregate(
            pending=_Count("id", filter=Q(status="pending")),
            scored=_Count("id", filter=Q(status="scored")),
            reconciled=_Count("id", filter=Q(status="reconciled")),
        )

        stats = {
            "total_students": student_agg["total"],
            "active_students": student_agg["active"],
            "total_examiners": User.objects.filter(role="examiner").count(),
            "total_procedures": Procedure.objects.count(),
            "pending_assessments": sp_agg["pending"],
            "scored_assessments": sp_agg["scored"],
            "reconciled_assessments": sp_agg["reconciled"],
            "total_programs": Program.objects.count(),
        }
        return Response(DashboardStatsSerializer(stats).data)


# ─────────────────────────────────────────────
# ADMIN: EXAMINERS
# ─────────────────────────────────────────────


# class ExaminerViewSet(viewsets.ModelViewSet):
#     queryset = User.objects.filter(role="examiner")
#     permission_classes = [IsAuthenticated, IsAdmin]

#     def get_serializer_class(self):
#         return UserCreateSerializer if self.action == "create" else UserSerializer

#     @action(detail=True, methods=["post"])
#     def toggle_active(self, request, pk=None):
#         user = self.get_object()
#         user.is_active = not user.is_active
#         user.save(update_fields=["is_active"])
#         return Response({"is_active": user.is_active})


class ExaminerViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated, IsAdmin]
    filter_backends = [filters.SearchFilter, filters.OrderingFilter]
    search_fields = ["username", "first_name", "last_name", "email"]
    ordering_fields = ["username", "first_name", "date_joined", "is_active"]
    ordering = ["username"]
 
    DEFAULT_IMPORT_PASSWORD = "Change123!"
 
    def get_queryset(self):
        qs = User.objects.filter(role="examiner")
        is_active = self.request.query_params.get("is_active")
        if is_active == "true":
            qs = qs.filter(is_active=True)
        elif is_active == "false":
            qs = qs.filter(is_active=False)
        return qs
 
    def get_serializer_class(self):
        return UserCreateSerializer if self.action == "create" else UserSerializer
 
    # ── Per-row toggle ────────────────────────────────────────────────────────
    @action(detail=True, methods=["post"])
    def toggle_active(self, request, pk=None):
        user = self.get_object()
        user.is_active = not user.is_active
        user.save(update_fields=["is_active"])
        return Response({"is_active": user.is_active})
 
    # ── Bulk delete ───────────────────────────────────────────────────────────
    @action(detail=False, methods=["post"], url_path="bulk-delete")
    @transaction.atomic
    def bulk_delete(self, request):
        ids = request.data.get("examiner_ids", [])
        if not ids or not isinstance(ids, list):
            return Response(
                {"error": "examiner_ids must be a non-empty list"}, status=400
            )
        count, _ = User.objects.filter(id__in=ids, role="examiner").delete()
        if count == 0:
            return Response(
                {"error": "No examiners found with the provided IDs"}, status=404
            )
        return Response(
            {
                "success": True,
                "deleted_count": count,
                "message": f"Successfully deleted {count} examiner(s)",
            }
        )
 
    # ── Bulk toggle active ────────────────────────────────────────────────────
    @action(detail=False, methods=["post"], url_path="bulk-toggle-active")
    @transaction.atomic
    def bulk_toggle_active(self, request):
        ids = request.data.get("examiner_ids", [])
        is_active = request.data.get("is_active")
 
        if not ids or not isinstance(ids, list):
            return Response(
                {"error": "examiner_ids must be a non-empty list"}, status=400
            )
        if is_active is None:
            return Response(
                {"error": "is_active (true/false) is required"}, status=400
            )
 
        count = User.objects.filter(id__in=ids, role="examiner").update(
            is_active=bool(is_active)
        )
        action_word = "activated" if is_active else "deactivated"
        return Response(
            {
                "success": True,
                "updated_count": count,
                "message": f"Successfully {action_word} {count} examiner(s)",
            }
        )
 
    # ── Import ────────────────────────────────────────────────────────────────
    @action(detail=False, methods=["post"], url_path="import")
    def import_examiners(self, request):
        if "file" not in request.FILES:
            return Response({"error": "No file provided"}, status=400)
 
        file = request.FILES["file"]
        ext = file.name.rsplit(".", 1)[-1].lower()
        if ext not in ("csv", "xlsx", "xls"):
            return Response(
                {"error": "Invalid file format. Use CSV or Excel."}, status=400
            )
        try:
            if ext == "csv":
                decoded = file.read().decode("utf-8").splitlines()
                rows = list(csv.DictReader(decoded))
            else:
                wb = load_workbook(file)
                ws = wb.active
                headers = [cell.value for cell in ws[1]]
                rows = [
                    dict(zip(headers, row))
                    for row in ws.iter_rows(min_row=2, values_only=True)
                    if any(row)
                ]
            return self._process_import(rows)
        except Exception as e:
            return Response({"error": str(e)}, status=400)
 
    @transaction.atomic
    def _process_import(self, rows):
        created = updated = errors = 0
        error_details = []
        success_details = []
 
        for row_num, row in enumerate(rows, start=2):
            try:
                username = str(row.get("Username") or "").strip()
                first_name = str(row.get("First Name") or "").strip()
                last_name = str(row.get("Last Name") or "").strip()
                email = str(row.get("Email") or "").strip()
                password = str(row.get("Password") or "").strip()
                status_raw = str(row.get("Status") or "Yes").strip().lower()
                is_active = status_raw in ("yes", "true", "1", "active")
 
                if not username:
                    error_details.append(
                        f"Row {row_num}: Username is required"
                    )
                    errors += 1
                    continue
 
                # Check for duplicate username in other roles
                conflict = User.objects.filter(username=username).exclude(
                    role="examiner"
                ).first()
                if conflict:
                    error_details.append(
                        f"Row {row_num}: Username '{username}' already exists "
                        f"as a {conflict.role}"
                    )
                    errors += 1
                    continue
 
                existing = User.objects.filter(
                    username=username, role="examiner"
                ).first()
 
                if existing:
                    # Update existing examiner (don't overwrite password unless provided)
                    existing.first_name = first_name or existing.first_name
                    existing.last_name = last_name or existing.last_name
                    existing.email = email or existing.email
                    existing.is_active = is_active
                    if password:
                        existing.set_password(password)
                    existing.save()
                    updated += 1
                else:
                    # Create new examiner
                    effective_password = password or self.DEFAULT_IMPORT_PASSWORD
                    User.objects.create_user(
                        username=username,
                        first_name=first_name,
                        last_name=last_name,
                        email=email,
                        password=effective_password,
                        role="examiner",
                        is_active=is_active,
                    )
                    created += 1
                    full = f"{first_name} {last_name}".strip() or username
                    success_details.append(
                        f"{full} (@{username})"
                        + ("" if password else " — default password assigned")
                    )
 
            except Exception as e:
                error_details.append(f"Row {row_num}: {e}")
                errors += 1
 
        return Response(
            {
                "success": True,
                "created": created,
                "updated": updated,
                "errors": errors,
                "error_details": error_details[:20],
                "success_details": success_details[:20],
            }
        )
 
    # ── Download import template ──────────────────────────────────────────────
    @action(detail=False, methods=["get"], url_path="template")
    def download_template(self, request):
        wb = Workbook()
        ws = wb.active
        ws.title = "Examiners Template"
 
        headers = ["Username", "First Name", "Last Name", "Email", "Password", "Status"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="4472C4", end_color="4472C4", fill_type="solid"
            )
 
        # Example rows
        ws.append(["jdoe", "John", "Doe", "jdoe@hospital.org", "SecurePass1!", "Yes"])
        ws.append(["asmith", "Alice", "Smith", "asmith@hospital.org", "", "Yes"])
 
        ws_inst = wb.create_sheet("Instructions")
        for row in [
            ["Import Instructions"],
            [""],
            ["Required columns:"],
            ["  Username (required), First Name, Last Name, Email, Password, Status"],
            [""],
            ["Notes:"],
            ["  - Username must be unique across all users."],
            ["  - Password is optional. If left blank, a default password is assigned."],
            ["  - Status accepts: Yes / No (defaults to Yes / Active)."],
            ["  - Existing examiners with the same username will be updated."],
        ]:
            ws_inst.append(row)
 
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value), default=12)
            ws.column_dimensions[col[0].column_letter].width = min(width + 4, 40)
 
        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        )
        response["Content-Disposition"] = (
            'attachment; filename="examiners_import_template.xlsx"'
        )
        wb.save(response)
        return response


# ─────────────────────────────────────────────
# ADMIN: STUDENTS
# ─────────────────────────────────────────────


class StudentViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated, IsAdmin]
    pagination_class = StudentPagination
    filter_backends = [
        DjangoFilterBackend,
        filters.SearchFilter,
        filters.OrderingFilter,
    ]
    filterset_class = StudentFilter
    search_fields = ["full_name", "index_number"]
    ordering_fields = ["full_name", "index_number", "level__number", "program__name"]
    ordering = ["level__number", "index_number"]

    def get_queryset(self):
        return Student.objects.select_related("program", "level").all()

    def get_serializer_class(self):
        if self.action in ("create", "update", "partial_update"):
            return StudentCreateUpdateSerializer
        return StudentSerializer

    def list(self, request, *args, **kwargs):
        export_format = request.query_params.get("export")
        if export_format:
            return self._handle_export(request, export_format)
        return super().list(request, *args, **kwargs)

    def _handle_export(self, request, export_format):
        students = self.filter_queryset(self.get_queryset()).filter(is_active=True)
        data = [
            {
                "index_number": s.index_number,
                "full_name": s.full_name,
                "program_name": s.program.name if s.program else "",
                "level": s.level.name if s.level else "",
                "is_active": "Yes" if s.is_active else "No",
            }
            for s in students
        ]
        if export_format == "csv":
            return self._export_csv(data)
        if export_format == "xlsx":
            return self._export_excel(data)
        if export_format == "pdf":
            return self._export_pdf(data)
        return Response({"error": "Invalid format"}, status=400)

    def _export_csv(self, data):
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = 'attachment; filename="students.csv"'
        writer = csv.writer(response)
        writer.writerow(["Index Number", "Full Name", "Program", "Level", "Status"])
        for item in data:
            writer.writerow(
                [
                    item["index_number"],
                    item["full_name"],
                    item["program_name"],
                    item["level"],
                    item["is_active"],
                ]
            )
        return response

    def _export_excel(self, data):
        wb = Workbook()
        ws = wb.active
        ws.title = "Students"
        headers = ["Index Number", "Full Name", "Program", "Level", "Status"]
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for item in data:
            ws.append(
                [
                    item["index_number"],
                    item["full_name"],
                    item["program_name"],
                    item["level"],
                    item["is_active"],
                ]
            )
        for column in ws.columns:
            width = max((len(str(c.value)) for c in column if c.value), default=10)
            ws.column_dimensions[column[0].column_letter].width = min(width + 2, 50)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="students.xlsx"'
        wb.save(response)
        return response

    def _export_pdf(self, data):
        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="students.pdf"'
        doc = SimpleDocTemplate(response, pagesize=landscape(letter))
        styles = getSampleStyleSheet()
        elements = [
            Paragraph("Students List", styles["Title"]),
            Paragraph("<br/><br/>", styles["Normal"]),
        ]
        table_data = [["Index Number", "Full Name", "Program", "Level", "Status"]]
        for item in data:
            table_data.append(
                [
                    item["index_number"],
                    item["full_name"],
                    item["program_name"],
                    item["level"],
                    item["is_active"],
                ]
            )
        table = Table(table_data)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                    ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, 0), 10),
                    ("BOTTOMPADDING", (0, 0), (-1, 0), 12),
                    ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
                    ("GRID", (0, 0), (-1, -1), 1, colors.black),
                    ("FONTSIZE", (0, 1), (-1, -1), 9),
                ]
            )
        )
        elements.append(table)
        doc.build(elements)
        return response

    @action(detail=True, methods=["post"])
    def toggle_active(self, request, pk=None):
        student = self.get_object()
        student.is_active = not student.is_active
        student.save(update_fields=["is_active"])
        return Response({"is_active": student.is_active})


class ImportStudentsView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def post(self, request):
        if "file" not in request.FILES:
            return Response({"error": "No file provided"}, status=400)
        file = request.FILES["file"]
        ext = file.name.rsplit(".", 1)[-1].lower()
        if ext not in ("csv", "xlsx", "xls"):
            return Response(
                {"error": "Invalid file format. Use CSV or Excel."}, status=400
            )
        try:
            if ext == "csv":
                return self._import_csv(file)
            return self._import_excel(file)
        except Exception as e:
            return Response({"error": str(e)}, status=400)

    def _import_csv(self, file):
        decoded = file.read().decode("utf-8").splitlines()
        return self._process_import(csv.DictReader(decoded))

    def _import_excel(self, file):
        wb = load_workbook(file)
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        data = [
            dict(zip(headers, row))
            for row in ws.iter_rows(min_row=2, values_only=True)
            if any(row)
        ]
        return self._process_import(data)

    @transaction.atomic
    def _process_import(self, data):
        created = updated = errors = 0
        error_details = []
        valid_levels = {"Level 100", "Level 200", "Level 300", "Level 400"}

        # Pre-load programs/levels into memory to avoid repeated queries
        programs = {p.name: p for p in Program.objects.all()}
        levels = {lv.name: lv for lv in Level.objects.all()}

        for row_num, row in enumerate(data, start=2):
            try:
                index_number = str(row.get("Index Number", "")).strip()
                full_name = str(row.get("Full Name", "")).strip()
                program_name = str(row.get("Program", "")).strip()
                level_str = str(row.get("Level", "Level 100")).strip()
                is_active = str(row.get("Status", "Yes")).strip().lower() in (
                    "yes",
                    "true",
                    "1",
                    "active",
                )

                if not index_number or not full_name or not program_name:
                    error_details.append(f"Row {row_num}: Missing required fields")
                    errors += 1
                    continue

                if level_str not in valid_levels:
                    error_details.append(
                        f"Row {row_num}: Invalid level '{level_str}'. "
                        "Must be Level 100, 200, 300, or 400"
                    )
                    errors += 1
                    continue

                if level_str not in levels:
                    levels[level_str], _ = Level.objects.get_or_create(name=level_str)
                if program_name not in programs:
                    programs[program_name], _ = Program.objects.get_or_create(
                        name=program_name
                    )

                _, was_created = Student.objects.update_or_create(
                    index_number=index_number,
                    defaults={
                        "full_name": full_name,
                        "program": programs[program_name],
                        "level": levels[level_str],
                        "is_active": is_active,
                    },
                )
                if was_created:
                    created += 1
                else:
                    updated += 1
            except Exception as e:
                error_details.append(f"Row {row_num}: {e}")
                errors += 1

        return Response(
            {
                "success": True,
                "created": created,
                "updated": updated,
                "errors": errors,
                "error_details": error_details[:10],
            }
        )


class DownloadStudentTemplateView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        wb = Workbook()
        ws = wb.active
        ws.title = "Students Template"
        ws.append(["Index Number", "Full Name", "Program", "Level", "Status"])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="4472C4", end_color="4472C4", fill_type="solid"
            )
        ws.append(
            ["L100-001", "John Doe", "Registered General Nursing", "Level 100", "Yes"]
        )
        ws.append(
            ["L200-002", "Jane Smith", "Public Health Nursing", "Level 200", "Yes"]
        )

        ws_inst = wb.create_sheet("Instructions")
        for row in [
            ["Import Instructions"],
            [""],
            ["Required columns:"],
            ["  Index Number, Full Name, Program (exact name), Level, Status (Yes/No)"],
            ["Level options: Level 100, Level 200, Level 300, Level 400"],
            ["Existing students (same Index Number) will be updated."],
        ]:
            ws_inst.append(row)

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = (
            'attachment; filename="students_import_template.xlsx"'
        )
        wb.save(response)
        return response


class BulkDeleteStudentsView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    @transaction.atomic
    def post(self, request):
        ids = request.data.get("student_ids", [])
        if not ids or not isinstance(ids, list):
            return Response(
                {"error": "student_ids must be a non-empty list"}, status=400
            )
        count, _ = Student.objects.filter(id__in=ids).delete()
        if count == 0:
            return Response(
                {"error": "No students found with provided IDs"}, status=404
            )
        return Response(
            {
                "success": True,
                "deleted_count": count,
                "message": f"Successfully deleted {count} student(s)",
            }
        )


# ─────────────────────────────────────────────
# ADMIN: PROCEDURES
# ─────────────────────────────────────────────


class ProcedureViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAuthenticated, IsAdmin]
    pagination_class = ProcedurePagination
    filter_backends = [
        DjangoFilterBackend,
        filters.SearchFilter,
    ]
    search_fields = ["name"]

    def get_queryset(self):
        qs = Procedure.objects.select_related("program").annotate(
            step_count=Count("steps")
        )
        program_id = self.request.query_params.get("program_id")
        if program_id and program_id != "all":
            qs = qs.filter(program_id=program_id)
        return qs

    def get_serializer_class(self):
        if self.action in ("create", "update", "partial_update"):
            return ProcedureCreateUpdateSerializer
        if self.action == "retrieve":
            return ProcedureDetailSerializer
        return ProcedureAdminListSerializer

    def list(self, request, *args, **kwargs):
        export_format = request.query_params.get("export")
        if export_format:
            return self._handle_export(request, export_format)
        return super().list(request, *args, **kwargs)

    def _handle_export(self, request, export_format):
        procedures = Procedure.objects.select_related("program").prefetch_related(
            "steps"
        )
        program_id = request.query_params.get("program_id")
        if program_id and program_id != "all":
            procedures = procedures.filter(program_id=program_id)

        if export_format == "excel":
            return self._export_excel(procedures)
        if export_format == "csv":
            return self._export_csv(procedures)
        if export_format == "pdf":
            return self._export_pdf(procedures)
        return Response({"error": "Invalid format"}, status=400)

    def _export_excel(self, procedures):
        wb = Workbook()
        ws_proc = wb.active
        ws_proc.title = "Procedures"
        headers_p = ["Name", "Program", "Total Score", "Steps Count"]
        ws_proc.append(headers_p)
        for cell in ws_proc[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="4472C4", end_color="4472C4", fill_type="solid"
            )
        for proc in procedures:
            ws_proc.append(
                [proc.name, proc.program.name, proc.total_score, proc.steps.count()]
            )

        ws_steps = wb.create_sheet("Procedure Steps")
        ws_steps.append(["Procedure Name", "Step Order", "Description"])
        for cell in ws_steps[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="70AD47", end_color="70AD47", fill_type="solid"
            )
        for proc in procedures:
            for step in proc.steps.all().order_by("step_order"):
                ws_steps.append([proc.name, step.step_order, step.description])

        for ws in (ws_proc, ws_steps):
            for col in ws.columns:
                width = max((len(str(c.value)) for c in col if c.value), default=10)
                ws.column_dimensions[col[0].column_letter].width = min(width + 2, 80)

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = (
            'attachment; filename="procedures_and_steps.xlsx"'
        )
        wb.save(response)
        return response

    def _export_csv(self, procedures):
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = (
            'attachment; filename="procedures_and_steps.csv"'
        )
        writer = csv.writer(response)
        writer.writerow(
            [
                "Procedure Name",
                "Program",
                "Total Score",
                "Step Order",
                "Step Description",
            ]
        )
        for proc in procedures:
            steps = list(proc.steps.order_by("step_order"))
            if steps:
                for step in steps:
                    writer.writerow(
                        [
                            proc.name,
                            proc.program.name,
                            proc.total_score,
                            step.step_order,
                            step.description,
                        ]
                    )
            else:
                writer.writerow(
                    [proc.name, proc.program.name, proc.total_score, "", ""]
                )
        return response

    def _export_pdf(self, procedures):
        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = (
            'attachment; filename="procedures_and_steps.pdf"'
        )
        doc = SimpleDocTemplate(response, pagesize=landscape(letter))
        elements = []
        styles = getSampleStyleSheet()
        elements.append(Paragraph("Procedures and Steps", styles["Title"]))
        elements.append(Spacer(1, 20))
        for proc in procedures:
            elements.append(
                Paragraph(
                    f"<b>{proc.name}</b> – {proc.program.name} (Total Score: {proc.total_score})",
                    styles["Heading2"],
                )
            )
            elements.append(Spacer(1, 8))
            steps = list(proc.steps.order_by("step_order"))
            if steps:
                table_data = [["Step", "Description"]] + [
                    [str(s.step_order), s.description] for s in steps
                ]
                t = Table(table_data, colWidths=[50, 450])
                t.setStyle(
                    TableStyle(
                        [
                            ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                            ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                            ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                            ("FONTSIZE", (0, 0), (-1, -1), 9),
                            ("BACKGROUND", (0, 1), (-1, -1), colors.beige),
                            ("GRID", (0, 0), (-1, -1), 1, colors.black),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ]
                    )
                )
                elements.append(t)
            else:
                elements.append(Paragraph("<i>No steps defined</i>", styles["Normal"]))
            elements.append(Spacer(1, 20))
        doc.build(elements)
        return response


class BulkDeleteProceduresView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    @transaction.atomic
    def post(self, request):
        ids = request.data.get("procedure_ids", [])
        if not ids or not isinstance(ids, list):
            return Response(
                {"error": "procedure_ids must be a non-empty list"}, status=400
            )
        count, _ = Procedure.objects.filter(id__in=ids).delete()
        if count == 0:
            return Response(
                {"error": "No procedures found with provided ID(s)"}, status=404
            )
        return Response(
            {
                "success": True,
                "deleted_count": count,
                "message": f"Successfully deleted {count} procedure(s)",
            }
        )


class ImportProceduresView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def post(self, request):
        if "file" not in request.FILES:
            return Response({"error": "No file provided"}, status=400)
        file = request.FILES["file"]
        ext = file.name.rsplit(".", 1)[-1].lower()
        if ext == "csv":
            return self._import_csv(file)
        if ext in ("xlsx", "xls"):
            return self._import_excel(file)
        return Response({"error": "Invalid file format. Use CSV or Excel."}, status=400)

    def _import_csv(self, file):
        try:
            decoded = file.read().decode("utf-8").splitlines()
        except UnicodeDecodeError:
            return Response(
                {"error": "File encoding error. Save as UTF-8."}, status=400
            )
        return self._process_csv_data(csv.DictReader(decoded))

    @transaction.atomic
    def _process_csv_data(self, reader):
        procs_created = procs_updated = steps_created = steps_updated = 0
        errors = []
        groups = {}
        for row_num, row in enumerate(reader, start=2):
            name = row.get("Procedure Name", "").strip()
            if not name:
                continue
            if name not in groups:
                groups[name] = {
                    "program_name": row.get("Program", "").strip(),
                    "total_score": row.get("Total Score", "").strip(),
                    "steps": [],
                }
            order_s = row.get("Step Order", "").strip()
            desc = row.get("Step Description", "").strip()
            if order_s and desc:
                try:
                    groups[name]["steps"].append(
                        {"order": int(order_s), "description": desc}
                    )
                except ValueError:
                    errors.append(f"Row {row_num}: Invalid step order '{order_s}'")

        programs = {p.name: p for p in Program.objects.all()}

        for name, data in groups.items():
            try:
                total_score = int(data["total_score"])
            except (ValueError, TypeError):
                errors.append(f"Procedure '{name}': Invalid total score")
                continue
            prog_name = data["program_name"]
            if prog_name:
                prog = programs.get(prog_name)
                if not prog:
                    errors.append(
                        f"Procedure '{name}': Program '{prog_name}' not found"
                    )
                    continue
                target_programs = [prog]
            else:
                target_programs = list(programs.values())
            for prog in target_programs:
                proc, created = Procedure.objects.update_or_create(
                    name=name, program=prog, defaults={"total_score": total_score}
                )
                if created:
                    procs_created += 1
                else:
                    procs_updated += 1
                for step in data["steps"]:
                    _, sc = ProcedureStep.objects.update_or_create(
                        procedure=proc,
                        step_order=step["order"],
                        defaults={"description": step["description"]},
                    )
                    if sc:
                        steps_created += 1
                    else:
                        steps_updated += 1
        return Response(
            {
                "success": True,
                "procedures_created": procs_created,
                "procedures_updated": procs_updated,
                "steps_created": steps_created,
                "steps_updated": steps_updated,
                "errors": len(errors),
                "error_details": errors[:20],
            }
        )

    def _import_excel(self, file):
        try:
            wb = load_workbook(file, data_only=True)
        except Exception as e:
            return Response({"error": f"Failed to read Excel file: {e}"}, status=400)
        if "Procedures" not in wb.sheetnames:
            return Response(
                {"error": 'Sheet "Procedures" not found in Excel file'}, status=400
            )

        procs_created = procs_updated = steps_created = steps_updated = 0
        errors = []
        procedures_dict = {}
        programs = {p.name: p for p in Program.objects.all()}

        with transaction.atomic():
            ws_proc = wb["Procedures"]
            for row_num, row in enumerate(
                ws_proc.iter_rows(min_row=2, values_only=True), start=2
            ):
                if not any(row):
                    continue
                try:
                    name = str(row[0]).strip() if row[0] else ""
                    prog_name = str(row[1]).strip() if row[1] else ""
                    try:
                        total_score = int(row[2]) if row[2] else 0
                    except (ValueError, TypeError):
                        errors.append(f"Procedures Row {row_num}: Invalid total score")
                        continue
                    if not name:
                        continue
                    if prog_name:
                        prog = programs.get(prog_name)
                        if not prog:
                            errors.append(
                                f"Procedures Row {row_num}: Program '{prog_name}' not found"
                            )
                            continue
                        target_progs = [prog]
                    else:
                        target_progs = list(programs.values())
                    for prog in target_progs:
                        proc, created = Procedure.objects.update_or_create(
                            name=name,
                            program=prog,
                            defaults={"total_score": total_score},
                        )
                        procedures_dict[(name, prog.name)] = proc
                        if created:
                            procs_created += 1
                        else:
                            procs_updated += 1
                except Exception as e:
                    errors.append(f"Procedures Row {row_num}: {e}")

            if "Procedure Steps" in wb.sheetnames:
                ws_steps = wb["Procedure Steps"]
                for row_num, row in enumerate(
                    ws_steps.iter_rows(min_row=2, values_only=True), start=2
                ):
                    if not any(row):
                        continue
                    try:
                        name = str(row[0]).strip() if row[0] else ""
                        try:
                            order = int(row[1]) if row[1] else 0
                        except (ValueError, TypeError):
                            errors.append(f"Steps Row {row_num}: Invalid step order")
                            continue
                        desc = str(row[2]).strip() if row[2] else ""
                        if not name or not desc:
                            continue
                        matching = [
                            proc
                            for (n, _), proc in procedures_dict.items()
                            if n == name
                        ]
                        if not matching:
                            matching = list(Procedure.objects.filter(name=name))
                        if not matching:
                            errors.append(
                                f"Steps Row {row_num}: Procedure '{name}' not found"
                            )
                            continue
                        for proc in matching:
                            _, sc = ProcedureStep.objects.update_or_create(
                                procedure=proc,
                                step_order=order,
                                defaults={"description": desc},
                            )
                            if sc:
                                steps_created += 1
                            else:
                                steps_updated += 1
                    except Exception as e:
                        errors.append(f"Steps Row {row_num}: {e}")

        return Response(
            {
                "success": True,
                "procedures_created": procs_created,
                "procedures_updated": procs_updated,
                "steps_created": steps_created,
                "steps_updated": steps_updated,
                "errors": len(errors),
                "error_details": errors[:20],
            }
        )


class DownloadProcedureTemplateView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        wb = Workbook()
        ws_proc = wb.active
        ws_proc.title = "Procedures"
        ws_proc.append(["Name", "Program", "Total Score"])
        for cell in ws_proc[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="4472C4", end_color="4472C4", fill_type="solid"
            )
        ws_proc.append(["Vital Signs Assessment", "Bachelor of Science in Nursing", 20])
        ws_proc.append(["IV Catheter Insertion", "Bachelor of Science in Nursing", 20])

        ws_steps = wb.create_sheet("Procedure Steps")
        ws_steps.append(["Procedure Name", "Step Order", "Description"])
        for cell in ws_steps[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="70AD47", end_color="70AD47", fill_type="solid"
            )
        ws_steps.append(
            [
                "Vital Signs Assessment",
                1,
                "Introduce yourself and explain the procedure",
            ]
        )
        ws_steps.append(["Vital Signs Assessment", 2, "Wash hands and put on gloves"])

        wb.create_sheet("Instructions").append(
            ["See column headers for required fields."]
        )

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = (
            'attachment; filename="procedures_import_template.xlsx"'
        )
        wb.save(response)
        return response


class ProcedureStepViewSet(viewsets.ModelViewSet):
    serializer_class = ProcedureStepCreateUpdateSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = ProcedureStep.objects.select_related("procedure")
        procedure_id = self.request.query_params.get("procedure_id")
        if procedure_id:
            qs = qs.filter(procedure_id=procedure_id)
        return qs


class ImportProcedureStepsView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def post(self, request, procedure_id):
        if "file" not in request.FILES:
            return Response({"error": "No file provided"}, status=400)
        file = request.FILES["file"]
        ext = file.name.rsplit(".", 1)[-1].lower()
        if ext not in ("csv", "xlsx", "xls"):
            return Response({"error": "Invalid file format."}, status=400)
        try:
            procedure = Procedure.objects.get(id=procedure_id)
        except Procedure.DoesNotExist:
            return Response({"error": "Procedure not found"}, status=404)
        try:
            if ext == "csv":
                decoded = file.read().decode("utf-8").splitlines()
                data = csv.DictReader(decoded)
            else:
                wb = load_workbook(file, data_only=True)
                ws = wb.active
                headers = [c.value for c in ws[1]]
                data = [
                    dict(zip(headers, row))
                    for row in ws.iter_rows(min_row=2, values_only=True)
                    if any(row)
                ]
            return self._process_import(data, procedure)
        except Exception as e:
            return Response({"error": str(e)}, status=400)

    @transaction.atomic
    def _process_import(self, data, procedure):
        created = updated = errors = 0
        error_details = []
        for row_num, row in enumerate(data, start=2):
            order_s = str(row.get("Step Order", "")).strip()
            desc = str(row.get("Description", "")).strip()
            if not order_s or not desc:
                error_details.append(
                    f"Row {row_num}: Missing step order or description"
                )
                errors += 1
                continue
            try:
                order = int(order_s)
            except ValueError:
                error_details.append(f"Row {row_num}: Invalid step order '{order_s}'")
                errors += 1
                continue
            _, was_created = ProcedureStep.objects.update_or_create(
                procedure=procedure,
                step_order=order,
                defaults={"description": desc},
            )
            if was_created:
                created += 1
            else:
                updated += 1
        return Response(
            {
                "success": True,
                "created": created,
                "updated": updated,
                "errors": errors,
                "error_details": error_details[:20],
            }
        )


class DownloadProcedureStepsTemplateView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request, procedure_id):
        try:
            procedure = Procedure.objects.get(id=procedure_id)
        except Procedure.DoesNotExist:
            return Response({"error": "Procedure not found"}, status=404)
        wb = Workbook()
        ws = wb.active
        ws.title = "Procedure Steps"
        ws.append(["Step Order", "Description"])
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                start_color="70AD47", end_color="70AD47", fill_type="solid"
            )
        steps = procedure.steps.order_by("step_order")
        if steps.exists():
            for step in steps:
                ws.append([step.step_order, step.description])
        else:
            ws.append([1, "Introduce yourself and explain the procedure"])
            ws.append([2, "Wash hands and put on gloves"])
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        safe_name = "".join(c for c in procedure.name if c.isalnum() or c in " _-")
        response["Content-Disposition"] = (
            f'attachment; filename="{safe_name}_steps_template.xlsx"'
        )
        wb.save(response)
        return response


# ─────────────────────────────────────────────
# ADMIN: GRADES
# ─────────────────────────────────────────────

SORT_FIELD_MAP = {
    "index_number": "index_number",
    "full_name": "full_name",
    "program": "program__name",
    "level": "level__number",
}


class StudentGradesView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        export_format = request.query_params.get("export")
        grade_filter = request.query_params.get("grade", "").strip()

        students = self._get_students(request)
        grades_data = self._build_grades_data(students)

        if grade_filter:
            grades_data = [
                g for g in grades_data if g["grade"].lower() == grade_filter.lower()
            ]

        # ───────── EXPORT ─────────
        if export_format:
            return self._handle_export(request, grades_data, grade_filter)

        sort_by = request.query_params.get("sort_by", "index_number")
        order = request.query_params.get("order", "asc")

        if sort_by in ("percentage", "total_score"):
            grades_data.sort(
                key=lambda x: x.get(sort_by, 0),
                reverse=(order == "desc"),
            )

        paginator = GradesPagination()
        page = paginator.paginate_queryset(grades_data, request)
        return paginator.get_paginated_response(page)

    # ------------------------------------------------------------------
    # Query builder
    # ------------------------------------------------------------------

    def _get_students(self, request):
        program_id = request.query_params.get("program_id")
        level_id = request.query_params.get("level_id")
        search = request.query_params.get("search", "").strip()

        procedure_score_sq = (
            StudentProcedure.objects.filter(student=OuterRef("pk"), status="reconciled")
            .values("student")
            .annotate(total=Sum("reconciled_scores__score"))
            .values("total")[:1]
        )
        procedure_max_sq = (
            StudentProcedure.objects.filter(student=OuterRef("pk"), status="reconciled")
            .values("student")
            .annotate(total=Sum("procedure__total_score"))
            .values("total")[:1]
        )
        procedure_count_sq = (
            StudentProcedure.objects.filter(student=OuterRef("pk"), status="reconciled")
            .values("student")
            .annotate(total=Count("id"))
            .values("total")[:1]
        )
        care_score_sq = (
            CarePlan.objects.filter(student=OuterRef("pk"))
            .values("student")
            .annotate(total=Sum("score"))
            .values("total")[:1]
        )
        care_max_sq = (
            CarePlan.objects.filter(student=OuterRef("pk"))
            .values("student")
            .annotate(total=Sum("max_score"))
            .values("total")[:1]
        )

        qs = (
            Student.objects.select_related("program", "level")
            .filter(is_active=True)
            .annotate(
                procedure_score=Coalesce(Subquery(procedure_score_sq), Value(0)),
                procedure_max_score=Coalesce(Subquery(procedure_max_sq), Value(0)),
                reconciled_count=Coalesce(Subquery(procedure_count_sq), Value(0)),
                care_plan_score=Coalesce(Subquery(care_score_sq), Value(0)),
                care_plan_max_score=Coalesce(Subquery(care_max_sq), Value(0)),
            )
        )

        if program_id:
            qs = qs.filter(program_id=program_id)
        if level_id and level_id != "all":
            qs = qs.filter(level_id=level_id)
        if search:
            qs = qs.filter(
                Q(full_name__icontains=search) | Q(index_number__icontains=search)
            )

        return qs

    def _build_grades_data(self, students):
        result = []
        for s in students:
            proc_score = s.procedure_score or 0
            proc_max = s.procedure_max_score or 0
            cp_score = s.care_plan_score or 0
            cp_max = s.care_plan_max_score or 0
            total = proc_score + cp_score
            max_score = proc_max + (cp_max if cp_score > 0 else 0)
            pct = round((total / max_score * 100), 1) if max_score > 0 else 0.0
            result.append(
                {
                    "student_id": s.id,
                    "index_number": s.index_number,
                    "full_name": s.full_name,
                    "program_name": s.program.name,
                    "program_id": s.program_id,
                    "level": s.level.name if s.level else "",
                    "level_id": s.level_id,
                    "procedure_score": round(proc_score, 2),
                    "procedure_max_score": proc_max,
                    "care_plan_score": cp_score,
                    "care_plan_max_score": cp_max,
                    "total_score": round(total, 2),
                    "max_score": max_score,
                    "percentage": pct,
                    "grade": self._calculate_grade(pct),
                    "reconciled_count": s.reconciled_count,
                    "care_plan_completed": cp_score > 0,
                }
            )
        return result

    def _calculate_grade(self, pct):
        if pct == 0:
            return "N/A"
        if pct >= 80:
            return "Distinction"
        if pct >= 70:
            return "Credit"
        if pct >= 60:
            return "Pass"
        return "Fail"

    # ------------------------------------------------------------------
    # Export handlers
    # ------------------------------------------------------------------
    def _handle_export(self, request, data, grade_filter):
        export_format = request.query_params.get("export")

        if export_format == "csv":
            return self._export_csv(data)

        if export_format == "excel":
            return self._export_excel(data)

        if export_format == "pdf":
            return self._export_pdf(
                data,
                request.query_params.get("program_id"),
                request.query_params.get("level_id"),
                grade_filter,
            )

        return Response({"error": "Invalid export format"}, status=400)

    EXPORT_HEADERS = [
        "Index Number",
        "Full Name",
        "Program",
        "Level",
        "Percentage (%)",
        "Grade",
    ]

    def _row(self, item):
        return [
            item["index_number"],
            item["full_name"],
            item["program_name"],
            item["level"],
            item["percentage"],
            item["grade"],
        ]

    def _export_csv(self, data):
        response = HttpResponse(content_type="text/csv")
        response["Content-Disposition"] = 'attachment; filename="student_grades.csv"'
        writer = csv.writer(response)
        writer.writerow(self.EXPORT_HEADERS)
        for item in data:
            writer.writerow(self._row(item))
        return response

    def _export_excel(self, data):
        wb = Workbook()
        ws = wb.active
        ws.title = "Student Grades"
        ws.append(self.EXPORT_HEADERS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for item in data:
            ws.append(self._row(item))
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(width + 2, 50)
        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response["Content-Disposition"] = 'attachment; filename="student_grades.xlsx"'
        wb.save(response)
        return response

    def _export_pdf(self, data, program_id=None, level_id=None, grade=None):
        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = 'attachment; filename="student_grades.pdf"'
        doc = SimpleDocTemplate(response, pagesize=landscape(letter))
        styles = getSampleStyleSheet()
        centered = ParagraphStyle(
            "Centered", parent=styles["Heading2"], alignment=TA_CENTER
        )
        elements = [Paragraph("STUDENT GRADES REPORT", styles["Title"])]
        if program_id:
            prog_name = (
                Program.objects.filter(id=program_id)
                .values_list("name", flat=True)
                .first()
            )
            if prog_name:
                elements.append(Paragraph(prog_name, centered))
        if level_id and level_id != "all":
            level_name = (
                Level.objects.filter(id=level_id).values_list("name", flat=True).first()
            )
            if level_name:
                elements.append(Paragraph(level_name, centered))
        if grade:
            elements.append(Paragraph(f"Grade: {grade}", centered))
        elements.append(Paragraph("<br/><br/>", styles["Normal"]))
        table_data = [self.EXPORT_HEADERS] + [self._row(item) for item in data]
        table_data[-len(data) :] = [
            [
                i["index_number"],
                i["full_name"],
                i["program_name"],
                i["level"],
                f"{i['percentage']}%",
                i["grade"],
            ]
            for i in data
        ]
        t = Table(table_data, repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                    ("ALIGN", (0, 0), (-1, -1), "LEFT"),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 1, colors.black),
                ]
            )
        )
        elements.append(t)
        doc.build(elements)
        return response


class GradeStatsView(APIView):
    permission_classes = [IsAuthenticated, IsAdmin]

    def get(self, request):
        view = StudentGradesView()

        students = view._get_students(request)
        data = view._build_grades_data(students)

        grade_filter = request.query_params.get("grade", "").strip().lower()

        # Apply same filtering logic
        if grade_filter:
            data = [g for g in data if g["grade"].lower() == grade_filter]

        total = len(data)
        completed = sum(1 for g in data if g["grade"] != "N/A")

        avg = round(sum(g["percentage"] for g in data) / total, 2) if total > 0 else 0

        grade_distribution = {}
        for g in data:
            grade_distribution[g["grade"]] = grade_distribution.get(g["grade"], 0) + 1

        complete_careplan = sum(1 for g in data if g["care_plan_completed"])

        return Response(
            {
                "total": total,
                "completed": completed,
                "average_percentage": avg,
                "grade_distribution": grade_distribution,
                "care_plan_completed": complete_careplan,
            }
        )
